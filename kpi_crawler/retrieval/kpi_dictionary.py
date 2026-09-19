"""Loader for the real, external KPI dictionary (xlsx) used by the retrieval baseline.

This is a source-of-record file the application does not own or generate. The
loader validates the sheet exists, the required identity/name/description
columns are present, and every row actually carries them — then fails loudly
and specifically rather than silently loading an empty or partial KPI set.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from ..errors import KPIDictionaryError

SHEET_NAME = "KPI_Dictionary"

# Required: the columns a KPI cannot be identified, named, or described without.
REQUIRED_COLUMNS = ("KPI_Code", "Variable_Name", "Definition")

# Optional "useful source fields" preserved as explicit columns on app.kpis.
# Everything else in the row (all 38 columns, including these) is preserved
# verbatim in source_row for full reproducibility.
OPTIONAL_COLUMNS = (
    "Domain",
    "Parameter",
    "Sub_Parameter",
    "Unit",
    "Data_Type",
    "Formula",
    "Primary_Source",
    "Benchmark_Direction",
    "Priority_12M",
)


@dataclass(frozen=True)
class KPIDefinition:
    kpi_code: str
    variable_name: str
    definition: str
    domain: str | None
    parameter: str | None
    sub_parameter: str | None
    unit: str | None
    data_type: str | None
    formula: str | None
    primary_source: str | None
    benchmark_direction: str | None
    priority_12m: str | None
    source_file: str
    source_row: dict[str, Any]

    @property
    def embedding_text(self) -> str:
        """The single, fixed KPI representation used for embedding: name + description."""
        return f"{self.variable_name}: {self.definition}"


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def load_kpi_dictionary(path: Path, *, sheet_name: str = SHEET_NAME) -> list[KPIDefinition]:
    """Load every KPI row from the real dictionary file.

    Raises KPIDictionaryError (never returns an empty/partial set silently) when
    the file is missing, unreadable, malformed, the expected sheet is absent, a
    required column is missing, or any row lacks its identity/name/description.
    """
    if not path.is_file():
        raise KPIDictionaryError(f"KPI dictionary file not found: {path}")

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several distinct exception types for bad files
        raise KPIDictionaryError(f"KPI dictionary file could not be read: {path} ({exc})") from exc

    try:
        if sheet_name not in workbook.sheetnames:
            raise KPIDictionaryError(
                f"KPI dictionary file {path} has no {sheet_name!r} sheet "
                f"(found: {workbook.sheetnames})"
            )
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise KPIDictionaryError(f"KPI dictionary sheet {sheet_name!r} in {path} is empty") from None

        header_names = [str(cell).strip() if cell is not None else None for cell in header]
        column_index = {name: idx for idx, name in enumerate(header_names) if name}

        missing_required = [name for name in REQUIRED_COLUMNS if name not in column_index]
        if missing_required:
            raise KPIDictionaryError(
                f"KPI dictionary sheet {sheet_name!r} in {path} is missing required "
                f"column(s): {missing_required} (found columns: {header_names})"
            )

        def cell(row: tuple[Any, ...], name: str) -> Any:
            idx = column_index.get(name)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        definitions: list[KPIDefinition] = []
        seen_codes: set[str] = set()
        for row_number, row in enumerate(rows, start=2):
            if all(value is None for value in row):
                continue  # a genuinely blank row (trailing sheet padding), not a KPI

            kpi_code = _clean(cell(row, "KPI_Code"))
            variable_name = _clean(cell(row, "Variable_Name"))
            definition = _clean(cell(row, "Definition"))
            missing_identity = [
                name
                for name, value in (
                    ("KPI_Code", kpi_code),
                    ("Variable_Name", variable_name),
                    ("Definition", definition),
                )
                if not value
            ]
            if missing_identity:
                raise KPIDictionaryError(
                    f"KPI dictionary row {row_number} in {path} is missing required "
                    f"value(s): {missing_identity}"
                )
            if kpi_code in seen_codes:
                raise KPIDictionaryError(
                    f"KPI dictionary in {path} has a duplicate KPI_Code: {kpi_code!r} (row {row_number})"
                )
            seen_codes.add(kpi_code)

            source_row = {name: row[idx] for name, idx in column_index.items() if idx < len(row)}

            definitions.append(
                KPIDefinition(
                    kpi_code=kpi_code,
                    variable_name=variable_name,
                    definition=definition,
                    domain=_clean(cell(row, "Domain")),
                    parameter=_clean(cell(row, "Parameter")),
                    sub_parameter=_clean(cell(row, "Sub_Parameter")),
                    unit=_clean(cell(row, "Unit")),
                    data_type=_clean(cell(row, "Data_Type")),
                    formula=_clean(cell(row, "Formula")),
                    primary_source=_clean(cell(row, "Primary_Source")),
                    benchmark_direction=_clean(cell(row, "Benchmark_Direction")),
                    priority_12m=_clean(cell(row, "Priority_12M")),
                    source_file=str(path),
                    source_row=source_row,
                )
            )
    finally:
        workbook.close()

    if not definitions:
        raise KPIDictionaryError(f"KPI dictionary sheet {sheet_name!r} in {path} contained no KPI rows")

    return definitions
